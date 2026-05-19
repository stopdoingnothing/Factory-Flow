"""Convert release-notes-sdn.md to a formatted Word document."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SPECS_DIR = Path(__file__).parent
INPUT_MD = SPECS_DIR / "release-notes-sdn.md"
OUTPUT_DOCX = SPECS_DIR / "release-notes-sdn.docx"

BRAND_BLUE = RGBColor(30, 90, 160)
DARK = RGBColor(30, 35, 45)
MID_GREY = RGBColor(140, 145, 155)
TABLE_HEADER_BG = "1E5AA0"
TABLE_ROW_ALT = "F0F2F6"


def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, start=120, bottom=60, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_border(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "CCCCCC")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def add_paragraph(doc, text="", style=None, bold=False, size=None, color=None, space_before=None, space_after=None, alignment=None):
    p = doc.add_paragraph(style=style)
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if alignment is not None:
        p.alignment = alignment
    if text:
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return p


def render_inline(para, text: str, default_color=None, default_size=None):
    """Render a line of text with **bold** and `code` inline markup."""
    pattern = re.compile(r'\*\*(.+?)\*\*|`(.+?)`')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = para.add_run(text[pos:m.start()])
            if default_color:
                run.font.color.rgb = default_color
            if default_size:
                run.font.size = Pt(default_size)
        if m.group(1):  # bold
            run = para.add_run(m.group(1))
            run.bold = True
            if default_color:
                run.font.color.rgb = default_color
            if default_size:
                run.font.size = Pt(default_size)
        elif m.group(2):  # code
            run = para.add_run(m.group(2))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(180, 50, 30)
        pos = m.end()
    if pos < len(text):
        run = para.add_run(text[pos:])
        if default_color:
            run.font.color.rgb = default_color
        if default_size:
            run.font.size = Pt(default_size)


def render_table(doc, rows):
    """Render a markdown pipe table. rows[0] is the header."""
    if not rows:
        return
    col_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_border(table)

    for r_idx, row_data in enumerate(rows):
        tr = table.rows[r_idx]
        for c_idx, cell_text in enumerate(row_data):
            cell = tr.cells[c_idx]
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_bg(cell, TABLE_HEADER_BG)
                p = cell.paragraphs[0]
                run = p.add_run(cell_text.strip())
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
            else:
                if r_idx % 2 == 0:
                    set_cell_bg(cell, TABLE_ROW_ALT)
                p = cell.paragraphs[0]
                render_inline(p, cell_text.strip(), default_size=9.5)

    doc.add_paragraph()


def parse_table_rows(lines, start):
    """Extract table data rows from markdown, skipping the separator row."""
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c for c in lines[i].strip().split("|") if c != ""]
        # Skip separator row (---|---)
        if not re.match(r'^[\s\-:]+$', cells[0]):
            rows.append(cells)
        i += 1
    return rows, i


def build_doc(md_text: str) -> Document:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    lines = md_text.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i]

        # H1
        if line.startswith("# ") and not line.startswith("## "):
            text = line[2:].strip()
            p = doc.add_heading(level=1)
            p.clear()
            run = p.add_run(text)
            run.font.size = Pt(22)
            run.font.color.rgb = BRAND_BLUE
            run.bold = True
            p.paragraph_format.space_after = Pt(6)
            i += 1

        # H2
        elif line.startswith("## ") and not line.startswith("### "):
            text = line[3:].strip()
            p = doc.add_heading(level=2)
            p.clear()
            run = p.add_run(text)
            run.font.size = Pt(15)
            run.font.color.rgb = BRAND_BLUE
            run.bold = True
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            i += 1

        # H3
        elif line.startswith("### "):
            text = line[4:].strip()
            p = doc.add_heading(level=3)
            p.clear()
            run = p.add_run(text)
            run.font.size = Pt(12)
            run.font.color.rgb = DARK
            run.bold = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            i += 1

        # Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "CCCCCC")
            pBdr.append(bottom)
            pPr.append(pBdr)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            i += 1

        # Table
        elif line.strip().startswith("|"):
            rows, i = parse_table_rows(lines, i)
            render_table(doc, rows)

        # Bullet list
        elif line.startswith("- "):
            text = line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            render_inline(p, text, default_color=DARK, default_size=10.5)
            i += 1

        # Metadata lines (key: value under H1)
        elif re.match(r'^\*\*(.+?)\*\*:', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            render_inline(p, line, default_color=MID_GREY, default_size=10)
            i += 1

        # Empty line
        elif line.strip() == "":
            i += 1

        # Normal paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            render_inline(p, line.strip(), default_color=DARK, default_size=10.5)
            i += 1

    return doc


def main():
    md_text = INPUT_MD.read_text(encoding="utf-8")
    doc = build_doc(md_text)
    doc.save(OUTPUT_DOCX)
    print(f"Saved: {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
