"""Convert markdown spec files to formatted Word documents using python-docx."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SPECS_DIR = Path(__file__).parent

MD_FILES = [
    "spec-aece-access-leave-and-hr-system.md",
    "spec-aece-payroll.md",
    "review-gap-analysis.md",
]

BRAND_BLUE = RGBColor(30, 90, 160)
DARK = RGBColor(30, 35, 45)
MID_GREY = RGBColor(140, 145, 155)
CODE_BG = RGBColor(240, 242, 246)


def set_cell_bg(cell, hex_color: str):
    """Set table cell background colour."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, start=100, bottom=60, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_run_bold(para, text: str, color=None):
    run = para.add_run(text)
    run.bold = True
    if color:
        run.font.color.rgb = color
    return run


def inline_runs(para, text: str):
    """Parse **bold**, *italic*, `code` inline and add runs."""
    pattern = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)")
    last = 0
    for m in pattern.finditer(text):
        # plain text before match
        if m.start() > last:
            para.add_run(sanitize(text[last:m.start()]))
        if m.group(0).startswith("**"):
            r = para.add_run(sanitize(m.group(2)))
            r.bold = True
        elif m.group(0).startswith("*"):
            r = para.add_run(sanitize(m.group(3)))
            r.italic = True
        else:
            r = para.add_run(sanitize(m.group(4)))
            r.font.name = "Courier New"
            r.font.size = Pt(9)
        last = m.end()
    if last < len(text):
        para.add_run(sanitize(text[last:]))


def sanitize(text: str) -> str:
    """Strip remaining markdown links and clean up text."""
    text = re.sub(r"\[(.+?)\]\(.*?\)", r"\1", text)
    return text


def strip_md(text: str) -> str:
    """Full strip of markdown formatting to plain text."""
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    text = re.sub(r"\[(.*?)\]\(.*?\)", r"\1", text)
    return text.strip()


def setup_styles(doc: Document):
    """Configure document-level styles."""
    # Normal
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = DARK
    style.paragraph_format.space_after = Pt(6)

    # Heading 1
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(255, 255, 255)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(8)

    # Heading 2
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = BRAND_BLUE
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)

    # Heading 3
    h3 = doc.styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = DARK
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(3)


def add_heading1(doc, text):
    p = doc.add_heading(strip_md(text), level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Blue background shading via paragraph shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "1E5AA0")
    pPr.append(shd)
    # Indent text slightly
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    return p


def add_heading2(doc, text):
    p = doc.add_heading(strip_md(text), level=2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_heading3(doc, text):
    p = doc.add_heading(strip_md(text), level=3)
    return p


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    inline_runs(p, text)
    return p


def add_bullet(doc, text, depth=0):
    style = "List Bullet" if depth == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.6 + depth * 0.5)
    inline_runs(p, text)
    return p


def add_numbered(doc, text, n):
    p = doc.add_paragraph(style="List Number")
    inline_runs(p, text)
    return p


def add_code_block(doc, code: str):
    for line in (code.splitlines() or [""]):
        p = doc.add_paragraph()
        p.style = doc.styles["Normal"]
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.right_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line or " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(50, 50, 80)
        # Cell-like shading on the paragraph
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F0F2F6")
        pPr.append(shd)
    doc.add_paragraph()


def add_table(doc, rows):
    headers = [strip_md(c.strip()) for c in rows[0].strip("|").split("|")]
    col_count = len(headers)
    data_rows = []
    for row in rows[1:]:
        cells = [strip_md(c.strip()) for c in row.strip("|").split("|")]
        while len(cells) < col_count:
            cells.append("")
        data_rows.append(cells[:col_count])

    table = doc.add_table(rows=1 + len(data_rows), cols=col_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Set column widths proportionally
    total_width = Inches(6.3)
    col_max = [max(len(h), max((len(r[ci]) for r in data_rows), default=0))
               for ci, h in enumerate(headers)]
    total_chars = sum(col_max) or 1
    col_widths = [total_width * (m / total_chars) for m in col_max]

    # Header row
    hdr_row = table.rows[0]
    for ci, (header, w) in enumerate(zip(headers, col_widths)):
        cell = hdr_row.cells[ci]
        cell.width = w
        set_cell_bg(cell, "1E5AA0")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.name = "Calibri"
        run.font.size = Pt(10)

    # Data rows
    for ri, row_data in enumerate(data_rows):
        tr = table.rows[ri + 1]
        bg = "F5F6F8" if ri % 2 == 0 else "FFFFFF"
        for ci, (cell_text, w) in enumerate(zip(row_data, col_widths)):
            cell = tr.cells[ci]
            cell.width = w
            set_cell_bg(cell, bg)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            inline_runs(p, cell_text)
            p.runs and setattr(p.runs[0].font, "size", Pt(10))

    doc.add_paragraph()


def add_cover(doc, title: str):
    """Add a cover page."""
    # Big title paragraph with blue shading
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(12)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "1E5AA0")
    pPr.append(shd)
    run = p.add_run(strip_md(title))
    run.bold = True
    run.font.size = Pt(28)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle line
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(6)
    pPr2 = sub._p.get_or_add_pPr()
    shd2 = OxmlElement("w:shd")
    shd2.set(qn("w:val"), "clear")
    shd2.set(qn("w:color"), "auto")
    shd2.set(qn("w:fill"), "1E5AA0")
    pPr2.append(shd2)
    r2 = sub.add_run("AEC Electronics (Pty) Ltd")
    r2.font.size = Pt(13)
    r2.font.name = "Calibri"
    r2.font.color.rgb = RGBColor(200, 220, 255)

    # Date bar
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_before = Pt(6)
    pPr3 = date_p._p.get_or_add_pPr()
    shd3 = OxmlElement("w:shd")
    shd3.set(qn("w:val"), "clear")
    shd3.set(qn("w:color"), "auto")
    shd3.set(qn("w:fill"), "E8EBF0")
    pPr3.append(shd3)
    r3 = date_p.add_run("April 2026")
    r3.font.size = Pt(11)
    r3.font.name = "Calibri"
    r3.font.color.rgb = MID_GREY

    doc.add_page_break()


def render_markdown(doc: Document, md_text: str):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        # Headings
        if line.startswith("### "):
            add_heading3(doc, line[4:].strip())

        elif line.startswith("## "):
            add_heading2(doc, line[3:].strip())

        elif line.startswith("# "):
            add_heading1(doc, line[2:].strip())

        # Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            p = doc.add_paragraph("_" * 80)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.color.rgb = RGBColor(180, 185, 195)
                run.font.size = Pt(6)

        # Fenced code block
        elif line.strip().startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, "\n".join(code_lines))

        # Table
        elif "|" in line and i + 1 < len(lines) and re.match(r"^\|[-| :]+\|", lines[i + 1]):
            table_lines = [line]
            i += 1  # skip separator row
            i += 1
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue

        # Unordered list
        elif re.match(r"^(\s*)[-*+] ", line):
            indent_base = len(re.match(r"^(\s*)", line).group(1))
            while i < len(lines) and re.match(r"^(\s*)[-*+] ", lines[i]):
                m = re.match(r"^(\s*)[-*+] (.*)", lines[i])
                depth = (len(m.group(1)) - indent_base) // 2
                add_bullet(doc, m.group(2).strip(), depth)
                i += 1
            continue

        # Numbered list
        elif re.match(r"^\d+\. ", line):
            n = 1
            while i < len(lines) and re.match(r"^\d+\. ", lines[i]):
                text = re.sub(r"^\d+\. ", "", lines[i]).strip()
                add_numbered(doc, text, n)
                n += 1
                i += 1
            continue

        # Blank line
        elif line.strip() == "":
            pass  # Word styles handle spacing naturally

        # Paragraph
        else:
            para_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and \
                  not lines[i].startswith("#") and \
                  "|" not in lines[i] and \
                  not re.match(r"^(\s*)[-*+] |\d+\. |```", lines[i]):
                para_lines.append(lines[i])
                i += 1
            add_paragraph(doc, " ".join(para_lines))
            continue

        i += 1


def md_title(path: Path) -> str:
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            return line[2:].strip()
    return path.stem.replace("-", " ").title()


def convert(md_path: Path):
    text = md_path.read_text(encoding="utf-8")
    title = md_title(md_path)

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    setup_styles(doc)
    add_cover(doc, title)
    render_markdown(doc, text)

    out = md_path.with_suffix(".docx")
    doc.save(str(out))
    print(f"  Created: {out.name}")


if __name__ == "__main__":
    print("Converting spec files to Word documents...")
    for name in MD_FILES:
        path = SPECS_DIR / name
        if path.exists():
            convert(path)
        else:
            print(f"  Skipped (not found): {name}")
    print("Done.")
